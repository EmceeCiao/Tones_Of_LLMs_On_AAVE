#!/usr/bin/env python3
"""
round2_similarity.py

Computes cosine (embedding) similarity and Jaccard similarity (token + bigram)
for every SAE/AAVE response pair in Round 2.

Round 2 responses live in the study log JSONL (final_prompt_responses.json has
the prompts for Round 2 but empty response fields, so we read from the log).

Output:
  Final_Dataset/round2_similarity.xlsx   — three sheets: Per Prompt, By Category, Summary
  Printed table to stdout

Usage:
    export OPENAI_API_KEY="your-key"
    python scripts/round2_similarity.py
"""

from __future__ import annotations

import json
import math
import os
import re
import statistics
from collections import defaultdict
from pathlib import Path

from openai import OpenAI
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT   = Path(__file__).resolve().parent.parent
LOG_PATH    = REPO_ROOT / "Final_Dataset" / "study_logs" / "round2_20260502T010128Z.jsonl"
OUTPUT_PATH = REPO_ROOT / "Final_Dataset" / "round2_similarity.xlsx"
EMBED_MODEL = "text-embedding-3-large"
BATCH_SIZE  = 64


# ---------------------------------------------------------------------------
# Data loading
# ---------------------------------------------------------------------------

def load_pairs(path: Path) -> list[dict]:
    raw: dict[int, dict] = {}
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            record = json.loads(line)
            if record.get("record_type") != "trial_result":
                continue
            if record.get("status") != "success":
                continue
            idx     = record["prompt_index"]
            dialect = record["dialect"]
            text    = (record.get("response_text") or "").strip()
            if not text:
                continue
            raw.setdefault(idx, {"prompt_index": idx, "category": record.get("category", "")})
            raw[idx][f"{dialect.lower()}_response"] = text

    pairs = []
    for idx in sorted(raw):
        entry = raw[idx]
        if "sae_response" not in entry or "aave_response" not in entry:
            continue
        pairs.append({
            "pair_id":       idx,
            "category":      entry["category"],
            "sae_response":  entry["sae_response"],
            "aave_response": entry["aave_response"],
        })
    return pairs


# ---------------------------------------------------------------------------
# Embedding similarity
# ---------------------------------------------------------------------------

def batched(items: list, size: int) -> list[list]:
    return [items[i : i + size] for i in range(0, len(items), size)]


def get_embeddings(client: OpenAI, texts: list[str]) -> list[list[float]]:
    vectors: list[list[float]] = []
    for batch in batched(texts, BATCH_SIZE):
        resp = client.embeddings.create(model=EMBED_MODEL, input=batch)
        vectors.extend(row.embedding for row in resp.data)
    return vectors


def cosine(a: list[float], b: list[float]) -> float:
    dot    = sum(x * y for x, y in zip(a, b))
    norm_a = math.sqrt(sum(x * x for x in a))
    norm_b = math.sqrt(sum(x * x for x in b))
    if norm_a == 0 or norm_b == 0:
        return 0.0
    return dot / (norm_a * norm_b)


# ---------------------------------------------------------------------------
# Jaccard similarity
# ---------------------------------------------------------------------------

def tokenize(text: str) -> list[str]:
    return [t for t in re.sub(r"[^a-z0-9\s]", " ", text.lower()).split() if t]


def bigrams(tokens: list[str]) -> list[tuple[str, str]]:
    return [(tokens[i], tokens[i + 1]) for i in range(len(tokens) - 1)]


def jaccard(a: set, b: set) -> float:
    if not a and not b:
        return 1.0
    union = a | b
    return len(a & b) / len(union) if union else 0.0


def jaccard_scores(sae: str, aave: str) -> tuple[float, float]:
    sae_tok  = tokenize(sae)
    aave_tok = tokenize(aave)
    token_j  = jaccard(set(sae_tok), set(aave_tok))
    bigram_j = jaccard(set(bigrams(sae_tok)), set(bigrams(aave_tok)))
    return token_j, bigram_j


# ---------------------------------------------------------------------------
# Output helpers
# ---------------------------------------------------------------------------

def fmt(v: float) -> str:
    return f"{v:.4f}"


def print_table(headers: list[str], rows: list[list], col_widths: list[int]) -> None:
    sep = "+" + "+".join("-" * (w + 2) for w in col_widths) + "+"
    def row_line(cells):
        parts = [f" {str(cell):<{w}} " for cell, w in zip(cells, col_widths)]
        return "|" + "|".join(parts) + "|"
    print(sep)
    print(row_line(headers))
    print(sep)
    for row in rows:
        print(row_line(row))
    print(sep)


# ---------------------------------------------------------------------------
# Excel writer
# ---------------------------------------------------------------------------

HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT  = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="D6E4F0")
SECTION_FONT = Font(bold=True)
CENTER       = Alignment(horizontal="center")


def style_header(cell) -> None:
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER


def auto_width(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)


def write_xlsx(
    path: Path,
    per_prompt_rows: list[dict],
    category_rows: list[dict],
    summary: dict,
) -> None:
    wb = Workbook()

    # ── Sheet 1: Per Prompt ──────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Per Prompt"
    headers = ["Pair ID", "Category", "Cosine Similarity", "Token Jaccard", "Bigram Jaccard",
               "SAE Chars", "AAVE Chars"]
    for col, h in enumerate(headers, start=1):
        style_header(ws1.cell(row=1, column=col, value=h))
    for r, row in enumerate(per_prompt_rows, start=2):
        ws1.cell(r, 1, row["pair_id"])
        ws1.cell(r, 2, row["category"])
        ws1.cell(r, 3, round(row["cosine_similarity"], 4))
        ws1.cell(r, 4, round(row["token_jaccard"], 4))
        ws1.cell(r, 5, round(row["bigram_jaccard"], 4))
        ws1.cell(r, 6, row["sae_chars"])
        ws1.cell(r, 7, row["aave_chars"])
    auto_width(ws1)

    # ── Sheet 2: By Category ─────────────────────────────────────────────
    ws2 = wb.create_sheet("By Category")
    cat_headers = ["Category", "n", "Mean Cosine", "Median Cosine",
                   "Mean Token Jaccard", "Median Token Jaccard",
                   "Mean Bigram Jaccard", "Median Bigram Jaccard"]
    for col, h in enumerate(cat_headers, start=1):
        style_header(ws2.cell(row=1, column=col, value=h))
    for r, row in enumerate(category_rows, start=2):
        ws2.cell(r, 1, row["category"])
        ws2.cell(r, 2, row["n"])
        ws2.cell(r, 3, round(row["mean_cosine"], 4))
        ws2.cell(r, 4, round(row["median_cosine"], 4))
        ws2.cell(r, 5, round(row["mean_token_j"], 4))
        ws2.cell(r, 6, round(row["median_token_j"], 4))
        ws2.cell(r, 7, round(row["mean_bigram_j"], 4))
        ws2.cell(r, 8, round(row["median_bigram_j"], 4))
    auto_width(ws2)

    # ── Sheet 3: Summary ─────────────────────────────────────────────────
    ws3 = wb.create_sheet("Summary")
    n = summary["n_pairs"]
    sum_data = [
        ("Metric",                  "Value"),
        ("Pairs analyzed",          n),
        ("Source log",              summary["source_log"]),
        ("Embedding model",         summary["embed_model"]),
        ("",                        ""),
        ("--- Cosine Similarity ---",""),
        ("Mean",                    round(summary["cosine_mean"],    4)),
        ("Median",                  round(summary["cosine_median"],  4)),
        ("Min",                     round(summary["cosine_min"],     4)),
        ("Max",                     round(summary["cosine_max"],     4)),
        ("Std dev",                 round(summary["cosine_stdev"],   4) if n > 1 else "n/a"),
        ("",                        ""),
        ("--- Token Jaccard ---",   ""),
        ("Mean",                    round(summary["token_j_mean"],   4)),
        ("Median",                  round(summary["token_j_median"], 4)),
        ("Min",                     round(summary["token_j_min"],    4)),
        ("Max",                     round(summary["token_j_max"],    4)),
        ("Std dev",                 round(summary["token_j_stdev"],  4) if n > 1 else "n/a"),
        ("",                        ""),
        ("--- Bigram Jaccard ---",  ""),
        ("Mean",                    round(summary["bigram_j_mean"],  4)),
        ("Median",                  round(summary["bigram_j_median"],4)),
        ("Min",                     round(summary["bigram_j_min"],   4)),
        ("Max",                     round(summary["bigram_j_max"],   4)),
        ("Std dev",                 round(summary["bigram_j_stdev"], 4) if n > 1 else "n/a"),
    ]
    for r, (label, value) in enumerate(sum_data, start=1):
        label_cell = ws3.cell(r, 1, label)
        ws3.cell(r, 2, value)
        if str(label).startswith("---"):
            label_cell.fill = SECTION_FILL
            label_cell.font = SECTION_FONT
        elif r == 1:
            style_header(ws3.cell(r, 1))
            style_header(ws3.cell(r, 2))
    auto_width(ws3)

    wb.save(path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    api_key = os.environ.get("OPENAI_API_KEY")
    if not api_key:
        raise SystemExit("Error: OPENAI_API_KEY environment variable is not set.")

    if not LOG_PATH.exists():
        raise SystemExit(f"Round 2 log not found: {LOG_PATH}")

    client = OpenAI(api_key=api_key)
    pairs  = load_pairs(LOG_PATH)
    if not pairs:
        raise SystemExit("No complete SAE+AAVE pairs found in the Round 2 log.")
    print(f"Loaded {len(pairs)} Round 2 pairs from {LOG_PATH.name}")

    # Embeddings
    print(f"Embedding {len(pairs) * 2} responses with {EMBED_MODEL}...")
    sae_vecs  = get_embeddings(client, [p["sae_response"]  for p in pairs])
    aave_vecs = get_embeddings(client, [p["aave_response"] for p in pairs])

    # Compute all scores
    per_prompt_rows: list[dict] = []
    for pair, sae_v, aave_v in zip(pairs, sae_vecs, aave_vecs):
        cos           = cosine(sae_v, aave_v)
        token_j, bi_j = jaccard_scores(pair["sae_response"], pair["aave_response"])
        per_prompt_rows.append({
            "pair_id":           pair["pair_id"],
            "category":          pair["category"],
            "cosine_similarity": cos,
            "token_jaccard":     token_j,
            "bigram_jaccard":    bi_j,
            "sae_chars":         len(pair["sae_response"]),
            "aave_chars":        len(pair["aave_response"]),
        })

    # Category breakdown
    by_cat: dict[str, list[dict]] = defaultdict(list)
    for row in per_prompt_rows:
        by_cat[row["category"]].append(row)

    category_rows: list[dict] = []
    for cat in sorted(by_cat):
        rows = by_cat[cat]
        cos  = [r["cosine_similarity"] for r in rows]
        tok  = [r["token_jaccard"]     for r in rows]
        bi   = [r["bigram_jaccard"]    for r in rows]
        category_rows.append({
            "category":        cat,
            "n":               len(rows),
            "mean_cosine":     statistics.mean(cos),
            "median_cosine":   statistics.median(cos),
            "mean_token_j":    statistics.mean(tok),
            "median_token_j":  statistics.median(tok),
            "mean_bigram_j":   statistics.mean(bi),
            "median_bigram_j": statistics.median(bi),
        })

    # Overall summary
    all_cos = [r["cosine_similarity"] for r in per_prompt_rows]
    all_tok = [r["token_jaccard"]     for r in per_prompt_rows]
    all_bi  = [r["bigram_jaccard"]    for r in per_prompt_rows]
    n       = len(per_prompt_rows)
    summary = {
        "n_pairs":         n,
        "source_log":      LOG_PATH.name,
        "embed_model":     EMBED_MODEL,
        "cosine_mean":     statistics.mean(all_cos),
        "cosine_median":   statistics.median(all_cos),
        "cosine_min":      min(all_cos),
        "cosine_max":      max(all_cos),
        "cosine_stdev":    statistics.stdev(all_cos)    if n > 1 else 0.0,
        "token_j_mean":    statistics.mean(all_tok),
        "token_j_median":  statistics.median(all_tok),
        "token_j_min":     min(all_tok),
        "token_j_max":     max(all_tok),
        "token_j_stdev":   statistics.stdev(all_tok)   if n > 1 else 0.0,
        "bigram_j_mean":   statistics.mean(all_bi),
        "bigram_j_median": statistics.median(all_bi),
        "bigram_j_min":    min(all_bi),
        "bigram_j_max":    max(all_bi),
        "bigram_j_stdev":  statistics.stdev(all_bi)    if n > 1 else 0.0,
    }

    # Print per-prompt table
    print()
    print("PER-PROMPT RESULTS")
    pp_headers = ["ID", "Category", "Cosine", "Token J", "Bigram J"]
    pp_table   = [
        [r["pair_id"], r["category"], fmt(r["cosine_similarity"]),
         fmt(r["token_jaccard"]), fmt(r["bigram_jaccard"])]
        for r in per_prompt_rows
    ]
    print_table(pp_headers, pp_table, [4, 12, 8, 8, 8])

    # Print category table
    print()
    print("BY CATEGORY")
    cat_headers = ["Category", "n", "Mean Cosine", "Mean Token J", "Mean Bigram J"]
    cat_table   = [
        [r["category"], r["n"], fmt(r["mean_cosine"]),
         fmt(r["mean_token_j"]), fmt(r["mean_bigram_j"])]
        for r in category_rows
    ]
    print_table(cat_headers, cat_table, [12, 4, 12, 13, 14])

    # Print overall summary
    print()
    print("OVERALL SUMMARY")
    sum_headers = ["Metric", "Mean", "Median", "Min", "Max", "Std Dev"]
    sum_table   = [
        ["Cosine",   fmt(summary["cosine_mean"]),   fmt(summary["cosine_median"]),
         fmt(summary["cosine_min"]),   fmt(summary["cosine_max"]),   fmt(summary["cosine_stdev"])],
        ["Token J",  fmt(summary["token_j_mean"]),  fmt(summary["token_j_median"]),
         fmt(summary["token_j_min"]),  fmt(summary["token_j_max"]),  fmt(summary["token_j_stdev"])],
        ["Bigram J", fmt(summary["bigram_j_mean"]), fmt(summary["bigram_j_median"]),
         fmt(summary["bigram_j_min"]), fmt(summary["bigram_j_max"]), fmt(summary["bigram_j_stdev"])],
    ]
    print_table(sum_headers, sum_table, [10, 8, 8, 8, 8, 8])

    # Save xlsx
    write_xlsx(OUTPUT_PATH, per_prompt_rows, category_rows, summary)
    print(f"\nResults saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
